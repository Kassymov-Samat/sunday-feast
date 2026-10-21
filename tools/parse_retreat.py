# -*- coding: utf-8 -*-
"""Разбор файла «20260502_Меню на Майский ретрит.xlsx»: блюда и продукты на одного гостя.

Дневные листы: блок «Завтрак»/«Обед» → строка «количество взрослых | | N» → блюда:
строка с номером в колонке A и названием в B, ниже — продукты (B — название, C — количество на всех).
Лист «Меню буфета»: заголовок блюда в B, рядом в C — на сколько порций.
"""
import re
from pathlib import Path

import openpyxl

SRC = Path(r"C:\projects\20260502_Меню на Майский ретрит.xlsx")
DAY_SHEETS = ["3 мая", "4 мая", "5 мая", "6 мая", "7 мая", "8 мая", "9 мая", "10 мая"]
SKIP_BUFFET = {"пицца", "напитки", "доп инвентарь"}  # пицца — с колбасой; напитки — газировка


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ".").strip())
    except (TypeError, ValueError):
        return None


def parse_days(wb):
    dishes = []
    for sh in DAY_SHEETS:
        ws = wb[sh]
        people, meal, cur = None, None, None
        for row in ws.iter_rows(values_only=True):
            a, b, c = (list(row) + [None] * 3)[:3]
            sa, sb = str(a or "").strip(), str(b or "").strip()
            if sa.lower() in ("завтрак", "обед", "ужин"):
                meal, cur = sa, None
                continue
            if sa.lower().startswith("количество взрослых"):
                people = num(c) or num(b)
                continue
            if sb.lower().startswith(("итого", "наименование")) or sa.lower().startswith(("повар", "№")):
                cur = None
                continue
            if num(a) is not None and sb and not num(c):  # строка блюда: «1 | Каша ... | (пусто или 0)»
                cur = {"name": sb, "day": sh, "meal": meal, "people": people, "items": []}
                dishes.append(cur)
                continue
            if cur is not None and not sa and sb and num(c):
                cur["items"].append((sb, num(c)))
    return dishes


def parse_buffet(wb):
    """Блок блюда: строка-заголовок (B — название, C — порций), затем «Наименование», затем продукты."""
    ws = wb["Меню буфета"]
    rows = [((list(r) + [None] * 3)[1], (list(r) + [None] * 3)[2]) for r in ws.iter_rows(values_only=True)]
    rows = [(str(b or "").strip(), c) for b, c in rows]
    dishes, cur = [], None
    for i, (sb, c) in enumerate(rows):
        if not sb:
            continue
        nxt = next((s for s, _ in rows[i + 1:i + 3] if s), "")
        if sb.lower().startswith("продукты на буфет"):
            cur = {"name": "Пирожки с картошкой, капустой и яблоками", "day": "буфет", "meal": "Буфет", "people": None, "items": []}
            dishes.append(cur)
        elif sb.lower() == "количество" and cur is not None:
            cur["people"] = num(c)
        elif nxt.lower() == "наименование":  # заголовок блюда: «пицца | 30», «Чебурки», «бурфи»
            cur = {"name": sb, "day": "буфет", "meal": "Буфет", "people": num(c), "items": []}
            dishes.append(cur)
        elif sb.lower() == "наименование" or sb.lower().startswith("итого"):
            continue
        elif cur is not None and num(c):
            cur["items"].append((sb, num(c)))
    out = []
    for d in dishes:
        if d["name"].lower() in SKIP_BUFFET or not d["items"]:
            continue
        if not d["people"]:
            d["people"] = 350  # чебурки, баларамы, бурфи — на буфет ретрита (≈350 человек)
        out.append(d)
    return out


def load():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    return parse_days(wb) + parse_buffet(wb), wb


if __name__ == "__main__":
    import collections, sys
    sys.stdout.reconfigure(encoding="utf-8")
    dishes, _ = load()
    print("блюд (с повторами):", len(dishes))
    names = collections.Counter()
    for d in dishes:
        for n, q in d["items"]:
            names[n.strip().lower()] += 1
    for d in dishes:
        print(f"[{d['day']} {d['meal']} {d['people']}] {d['name']}: " + "; ".join(f"{n}={q}" for n, q in d["items"]))
    print("\n=== уникальные продукты:", len(names))
    for n, k in sorted(names.items()):
        print(f"{k:3} {n}")
